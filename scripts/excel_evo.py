import pickle
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
B = ROOT / 'build'

E = pickle.load(open(B / 'evo.pkl', 'rb'))
P = ROOT / 'outputs' / 'Base_sport_feminin_Val_d_Oise.xlsx'
wb = load_workbook(P)
FONT = 'Arial'
HFILL = PatternFill('solid', start_color='1F3A5F')
HFONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)
BFONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, bold=True, size=10)
TFONT = Font(name=FONT, bold=True, size=14, color='1F3A5F')
NFONT = Font(name=FONT, italic=True, size=9, color='555555')
PCT = '0.0%;-0.0%;"-"'
INT = '#,##0;-#,##0;"-"'
DEC = '+0.0;-0.0;"0,0"'

for name in ('Disciplines_5_saisons', 'Territoires_2019_2025'):
    if name in wb.sheetnames:
        del wb[name]


def hdr(ws, row, cols, widths):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c)
        x.font, x.fill = HFONT, HFILL
        x.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 44
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=2)


def put(ws, r, c, v, fmt=None, font=BFONT):
    if isinstance(v, (np.floating, float)) and np.isnan(v):
        v = None
    elif isinstance(v, np.integer):
        v = int(v)
    elif isinstance(v, np.floating):
        v = float(v)
    x = ws.cell(row=r, column=c, value=v)
    x.font = font
    if fmt:
        x.number_format = fmt
    return x


# ============ Disciplines 5 saisons
ws = wb.create_sheet('Disciplines_5_saisons', index=wb.sheetnames.index('Comites') + 1)
ws['A1'] = 'Part des femmes par discipline, 2020/21 à 2025/26'
ws['A1'].font = TFONT
ws['A2'] = ('Sources : enquête « Le sport féminin en Val d\'Oise » (2020/21) ; comités départementaux, onglet Bilan (2022/23 à 2025/26). '
            'La saison 2021/22 ne renseigne pas les féminines. 2020/21 est une saison perturbée par la crise sanitaire.')
ws['A2'].font = NFONT
SEAS = ['2020/2021', '2022/2023', '2023/2024', '2024/2025', '2025/2026']
cols = ['Discipline'] + sum([[f'Licencié·es {s}', f'Féminines {s}', f'Part F {s}'] for s in SEAS], []) + \
       ['Écart 2020/21 → 2025/26 (points)', 'Présente les 5 saisons']
hdr(ws, 4, cols, [26] + [11] * (len(cols) - 3) + [14, 12])
wide = E['wide']
names = sorted(wide.index, key=lambda s: s.lower())
r = 5
for n in names:
    put(ws, r, 1, n)
    for j, s in enumerate(SEAS):
        cl, cf, cp = 2 + 3 * j, 3 + 3 * j, 4 + 3 * j
        lv = wide.loc[n, ('licencies', s)] if ('licencies', s) in wide.columns else np.nan
        fv = wide.loc[n, ('feminines', s)] if ('feminines', s) in wide.columns else np.nan
        pv = wide.loc[n, ('p', s)] if ('p', s) in wide.columns else np.nan
        put(ws, r, cl, lv, INT)
        # write F only when valid (F <= total)
        put(ws, r, cf, fv if not np.isnan(pv) else None, INT)
        put(ws, r, cp, f'=IF(AND(ISNUMBER({L(cl)}{r}),ISNUMBER({L(cf)}{r}),N({L(cl)}{r})>0),{L(cf)}{r}/{L(cl)}{r},"")', PCT)
    cA, cB = L(4), L(4 + 3 * 4)
    put(ws, r, 17, f'=IF(AND(ISNUMBER({cA}{r}),ISNUMBER({cB}{r})),({cB}{r}-{cA}{r})*100,"")', DEC)
    put(ws, r, 18, 'oui' if n in E['okall'] else 'non')
    r += 1
last = r - 1
t = Table(displayName='T_Disc5', ref=f'A4:{L(len(cols))}{last}')
t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
ws.add_table(t)
r = last + 2
put(ws, r, 1, 'Total des disciplines présentes les 5 saisons', font=BOLD)
for j, s in enumerate(SEAS):
    cl, cf, cp = 2 + 3 * j, 3 + 3 * j, 4 + 3 * j
    put(ws, r, cl, f'=SUMIFS({L(cl)}5:{L(cl)}{last},$R$5:$R${last},"oui")', INT, BOLD)
    put(ws, r, cf, f'=SUMIFS({L(cf)}5:{L(cf)}{last},$R$5:$R${last},"oui")', INT, BOLD)
    put(ws, r, cp, f'={L(cf)}{r}/{L(cl)}{r}', PCT, BOLD)
put(ws, r, 17, f'=({L(16)}{r}-{L(4)}{r})*100', DEC, BOLD)
put(ws, r + 1, 1, f'{len(E["okall"])} disciplines renseignent licencié·es et féminines aux 5 saisons. '
                  'Taekwondo 2020/21 est exclu : féminines (543) supérieures au total (218).', font=NFONT)

# ============ Territoires 2019-2025
a = E['assos']
ws = wb.create_sheet('Territoires_2019_2025', index=wb.sheetnames.index('Evolution') + 1)
ws['A1'] = 'Associations civiles aidées par territoire, 2019 à 2025'
ws['A1'].font = TFONT
ws['A2'] = (f'Tous sexes confondus (pas de répartition F/H avant 2024). « Panel » = les {E["n7"]} associations présentes les 7 années, '
            'pour suivre les effectifs sans l\'effet des entrées et sorties de la liste. Volume = licencié·es jusqu\'en 2023, adhérent·es ensuite.')
ws['A2'].font = NFONT
YEARS = list(range(2019, 2026))
row = 4
for lvl, lab in (('canton', 'Canton'), ('epci', 'Intercommunalité')):
    n = a.groupby([lvl, 'annee']).k.size().unstack()
    v = a.groupby([lvl, 'annee']).vol.sum().unstack()
    p = a[a.constante].groupby([lvl, 'annee']).vol.sum().unstack()
    cols = [lab] + [f'Asso. {y}' for y in YEARS] + [f'Volume {y}' for y in YEARS] + [f'Panel {y}' for y in YEARS] + \
           ['Évol. asso. 2019→2025', 'Évol. panel 2019→2022', 'Évol. panel 2019→2025']
    put(ws, row, 1, f'Par {lab.lower()}', font=Font(name=FONT, bold=True, size=12, color='1F3A5F'))
    row += 1
    h = row
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=h, column=j, value=c)
        x.font, x.fill = HFONT, HFILL
        x.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[h].height = 32
    keys = sorted(n.index)
    for i, k in enumerate(keys, h + 1):
        put(ws, i, 1, k)
        for j, y in enumerate(YEARS):
            put(ws, i, 2 + j, n.loc[k].get(y, 0), INT)
            put(ws, i, 9 + j, v.loc[k].get(y, 0), INT)
            put(ws, i, 16 + j, p.loc[k].get(y, 0) if k in p.index else 0, INT)
        put(ws, i, 23, f'=H{i}-B{i}', '+0;-0;0')
        put(ws, i, 24, f'=IF(P{i}>0,S{i}/P{i}-1,"")', PCT)
        put(ws, i, 25, f'=IF(P{i}>0,V{i}/P{i}-1,"")', PCT)
    lastr = h + len(keys)
    tr = lastr + 1
    put(ws, tr, 1, 'Total Val d\'Oise', font=BOLD)
    for j in range(2, 23):
        put(ws, tr, j, f'=SUM({L(j)}{h + 1}:{L(j)}{lastr})', INT, BOLD)
    put(ws, tr, 23, f'=H{tr}-B{tr}', '+0;-0;0', BOLD)
    put(ws, tr, 24, f'=S{tr}/P{tr}-1', PCT, BOLD)
    put(ws, tr, 25, f'=V{tr}/P{tr}-1', PCT, BOLD)
    row = tr + 3
ws.column_dimensions['A'].width = 36
for j in range(2, 26):
    ws.column_dimensions[L(j)].width = 10
ws.freeze_panes = 'B4'
put(ws, row, 1, 'Anomalie connue : en 2022-2023, le CMG Club multisports de Garges-lès-Gonesse déclare 370 puis 37 licencié·es (1 830 en 2021, 1 330 en 2024).', font=NFONT)
put(ws, row + 1, 1, 'Les associations sont rapprochées d\'une année à l\'autre par leur nom. Une association renommée compte comme une sortie puis une entrée.', font=NFONT)

# Lisez-moi additions
lm = wb['Lisez-moi']
nr = lm.max_row + 2
extra = [('Ajouts 22/09/2026', ''),
         ('Disciplines_5_saisons', 'Part des femmes par discipline : 2020/21 (enquête) puis 2022/23 à 2025/26 (comités). Total calculé sur les disciplines présentes les 5 saisons.'),
         ('Territoires_2019_2025', 'Nombre d\'associations aidées, volume de licencié·es / adhérent·es et panel d\'associations présentes les 7 années, par canton et intercommunalité.')]
for i, (x, y) in enumerate(extra, nr):
    lm.cell(row=i, column=1, value=x).font = BOLD
    c = lm.cell(row=i, column=2, value=y)
    c.font = BFONT
    c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(P)
print('ok')
